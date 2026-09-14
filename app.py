
import streamlit as st
import pandas as pd
import openpyxl
import io

st.set_page_config(
    page_title="EIB Assistant – RSU Tax Formatter",
    page_icon="🧾",
    layout="wide",
)

st.markdown("""
<style>
#MainMenu, footer { visibility: hidden; }
.block-container { padding-top: 0 !important; padding-bottom: 2rem !important; }

/* Header */
.eib-header {
  background: linear-gradient(135deg, #1a1a2e 0%, #16213e 50%, #0f3460 100%);
  color: white; padding: 18px 32px; display: flex; align-items: baseline;
  gap: 14px; box-shadow: 0 2px 12px rgba(0,0,0,0.3);
  margin: -1rem -1rem 1.5rem -1rem;
}
.eib-logo { font-size: 1.6rem; font-weight: 700; letter-spacing: -0.5px; }
.eib-sub  { font-size: 0.9rem; color: #94b4e4; font-weight: 400; }

/* Cards */
.eib-card {
  background: white; border-radius: 14px; padding: 24px 28px;
  margin-bottom: 1rem; box-shadow: 0 2px 8px rgba(0,0,0,0.07);
  position: relative;
}
.eib-card h3 { font-size: 1.1rem; font-weight: 700; margin-bottom: 4px; color: #1a1a2e; }
.eib-card p  { color: #555; font-size: 0.88rem; margin: 0; }
.badge {
  display: inline-block; background: #e8f0fe; color: #1967d2;
  font-size: 0.7rem; font-weight: 700; padding: 2px 9px;
  border-radius: 20px; text-transform: uppercase; letter-spacing: 0.5px;
  margin-bottom: 6px;
}
.row-badge {
  display: inline-block; background: #f0f4ff; color: #1967d2;
  font-size: 0.8rem; font-weight: 700; padding: 3px 12px;
  border-radius: 20px; border: 1px solid #c8d8f8; margin-left: 8px;
  vertical-align: middle;
}

/* Legend */
.legend { display: flex; gap: 16px; flex-wrap: wrap; margin-bottom: 12px; }
.leg { display: flex; align-items: center; gap: 6px; font-size: 0.78rem; color: #555; }
.dot { width: 12px; height: 12px; border-radius: 3px; display: inline-block; flex-shrink: 0; }
.dot-tv { background:#c6e8ce; border:1px solid #a8d5b5; }
.dot-nv { background:#c2d8fb; border:1px solid #a0c0f5; }
.dot-ws { background:#fad67d; border:1px solid #f0bc40; }
.dot-sb { background:#fce4c3; border:1px solid #f0c890; }

/* Table */
.tbl-wrap { overflow-x: auto; margin-top: 8px; }
.eib-tbl { border-collapse: collapse; width: 100%; font-size: 0.82rem; }
.eib-tbl th {
  background: #f0f2f5; color: #444; font-weight: 700; padding: 8px 10px;
  text-align: left; border-bottom: 2px solid #d1d5db; white-space: nowrap;
  font-size: 0.74rem; text-transform: uppercase; letter-spacing: 0.3px;
}
.eib-tbl td { padding: 7px 10px; border-bottom: 1px solid #f0f2f5; white-space: nowrap; color: #1a1a2e; }
.eib-tbl tr.tv td { background: #e6f4ea; }
.eib-tbl tr.nv td { background: #e8f0fe; }
.eib-tbl tr.ws td { background: #fce8b2; }
.eib-tbl tr.sb td { background: #fdf0e0; }
.eib-tbl tr.dv td { border-top: 2.5px solid #c0cfe8; padding: 0; height: 3px; }
.eib-tbl td.num { text-align: right; font-variant-numeric: tabular-nums; }
</style>

<div class="eib-header">
  <span class="eib-logo">🧾 EIB Assistant</span>
  <span class="eib-sub">RSU Tax Formatter</span>
</div>
""", unsafe_allow_html=True)

# ── Constants ─────────────────────────────────────────────
STATE_MAP = {
    'AL':'01','AK':'02','AZ':'04','AR':'05','CA':'06','CO':'08','CT':'09',
    'DE':'10','DC':'11','FL':'12','GA':'13','GU':'66','HI':'15','ID':'16',
    'IL':'17','IN':'18','IA':'19','KS':'20','KY':'21','LA':'22','ME':'23',
    'MD':'24','MA':'25','MI':'26','MN':'27','MS':'28','MO':'29','MT':'30',
    'NE':'31','NV':'32','NH':'33','NJ':'34','NM':'35','NY':'36','NC':'37',
    'ND':'38','OH':'39','OK':'40','OR':'41','PA':'42','PR':'72','RI':'44',
    'SC':'45','SD':'46','TN':'47','TX':'48','UT':'49','VT':'50','VA':'51',
    'VI':'78','WA':'53','WV':'54','WI':'55','WY':'56',
}

TOTAL_COLS = 134
C = dict(spreadsheetKey=1,batchId=3,paymentId=4,employee=5,paymentDate=6,
         periodDate=7,payPriority=8,runCategory=9,resultType=11,reason=13,
         rowId=39,earning=40,deduction=41,amount=43,stateAuthority=57,
         subRowId=85,relatedCalc=86,value=87)

EIB_HDR = ['Fields','Spreadsheet Key*','Payroll Off-cycle Payment','Batch ID','Payment ID*',
    'Employee*','Payment Date*','Period Date*','Payment Priority*','Run Category','Pay Group',
    'Result Type*','Replacement','Reason*','Use Supplemental Tax Rate','Override Payment to Check',
    'Take Additional Withholding','Include Retro Differences in Payment','Load or Refresh Input',
    'Row ID*','Tax Frequency Value','Tax Frequency Period','Third Party Sick Pay','Net Amount',
    'Check Number','Bank Account','Company','Region','Location','Cost Center','Job Profile',
    'State  Work ','State  Resident ','County  Work ','County  Resident ','City  Work ',
    'City  Resident ','School District  Resident ','Payroll Reference Number','Row ID*',
    'Earning*','Deduction*','Position','Amount','Hours','Rate','Adjustment','Reference Date',
    'Currency','Location','Region','Job Profile','Cost Center','Project','Project Phase',
    'Project Task','Withholding Order Case','State Authority','County Authority','City Authority',
    'School District Authority','Custom Worktag 01','Custom Worktag 02','Custom Worktag 03',
    'Custom Worktag 04','Custom Worktag 05','Fund','Grant','Gift','Program','Business Unit',
    'Object Class','Custom Organization+','Custom Worktag 06','Custom Worktag 07',
    'Custom Worktag 08','Custom Worktag 09','Custom Worktag 10','Custom Worktag 11',
    'Custom Worktag 12','Custom Worktag 13','Custom Worktag 14','Custom Worktag 15',
    'Local Other Tax Authority','NI Category','Row ID*','Related Calculation','Value','Company']

DISPLAY = [('Spr. Key',C['spreadsheetKey']),('Batch ID',C['batchId']),
    ('Payment ID',C['paymentId']),('Employee',C['employee']),
    ('Pay Date',C['paymentDate']),('Period Date',C['periodDate']),
    ('Priority',C['payPriority']),('Run Cat.',C['runCategory']),
    ('Result Type',C['resultType']),('Reason',C['reason']),
    ('Row ID',C['rowId']),('Earning',C['earning']),('Deduction',C['deduction']),
    ('Amount',C['amount']),('State Auth',C['stateAuthority']),
    ('Sub Row ID',C['subRowId']),('Rel. Calc',C['relatedCalc']),('Value',C['value'])]

COL_KEYS = dict(
    empId        = ['employee id'],
    txnId        = ['transaction id'],
    name         = ['name'],
    stateCode    = ['payroll authority tax code','state code','region code'],
    totalDist    = ['total distribution/exercise gross gain','reportable income'],
    totalWith    = ['total employee withholding (actual & hypo taxes)','country total actual tax ee liability'],
    regionTax    = ['region tax employee liability'],
    incomeRegion = ['income for region tax pre gross up'],
)

# ── Helpers ───────────────────────────────────────────────
def resolve(raw):
    if not raw and raw != 0: return ''
    s = str(raw).strip().upper()
    if s in STATE_MAP: return STATE_MAP[s]
    try: return str(int(s)).zfill(2)
    except: return s

def fcol(headers, keys):
    for key in keys:
        for i,h in enumerate(headers):
            if h and key in str(h).lower().replace('\n',' '): return i
    return -1

def to_num(v):
    try: return float(v) if v not in (None,'') else None
    except: return None

def empty(): return [None]*TOTAL_COLS

def fmt(v):
    if v is None: return ''
    if isinstance(v,float): return f"{v:,.2f}"
    return str(v)

# ── Parse ─────────────────────────────────────────────────
def parse_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    hi, headers = -1, []
    for i,row in enumerate(rows[:5]):
        j = ' '.join(str(c).lower() for c in row if c)
        if 'employee' in j or 'transaction' in j:
            hi = i; headers = [str(c) if c else '' for c in row]; break
    if hi < 0: return []
    cE=fcol(headers,COL_KEYS['empId']); cT=fcol(headers,COL_KEYS['txnId'])
    cN=fcol(headers,COL_KEYS['name']); cS=fcol(headers,COL_KEYS['stateCode'])
    cD=fcol(headers,COL_KEYS['totalDist']); cW=fcol(headers,COL_KEYS['totalWith'])
    cR=fcol(headers,COL_KEYS['regionTax']); cI=fcol(headers,COL_KEYS['incomeRegion'])
    if cE<0 or cT<0: return []
    groups,order,le,lt = {},[],None,None
    for row in rows[hi+1:]:
        e = str(row[cE]).strip() if cE>=0 and row[cE] else None
        t = str(row[cT]).strip() if cT>=0 and row[cT] else None
        if not e and not t and all(v is None for v in row): continue
        if not e: e=le
        if not t: t=lt
        if not e or not t: continue
        le,lt = e,t; key=f"{e}|{t}"
        if key not in groups:
            groups[key]={'empId':e,'txnId':str(t),'name':'','rows':[]}
            order.append(key)
        if cN>=0 and row[cN] and not groups[key]['name']:
            groups[key]['name']=str(row[cN])
        sc = row[cS] if cS>=0 else None
        groups[key]['rows'].append({
            'stateCode':resolve(sc),
            'totalDist':to_num(row[cD] if cD>=0 else None),
            'totalWith':to_num(row[cW] if cW>=0 else None),
            'regionTax':to_num(row[cR] if cR>=0 else None),
            'incomeRegion':to_num(row[cI] if cI>=0 else None),
        })
    return [groups[k] for k in order]

# ── EIB generation ────────────────────────────────────────
def base(s, key, emp):
    r=empty()
    r[C['spreadsheetKey']]=key; r[C['batchId']]=s['batchId']
    r[C['paymentId']]='="Trailing2State"&D&F'; r[C['employee']]=emp
    r[C['paymentDate']]=s['paymentDate']; r[C['periodDate']]=s['periodDate']
    r[C['payPriority']]=int(s['payPriority']); r[C['runCategory']]=s['runCategory']
    r[C['resultType']]=s['resultType']; r[C['reason']]=s['reason']
    return r

def gen(records, s):
    out=[]
    for rec in records:
        primary = next((r for r in rec['rows'] if r['totalDist'] is not None), rec['rows'][0])
        key=rec['txnId']; line=1
        r=base(s,key,rec['empId']); r[C['rowId']]=line; line+=1
        r[C['earning']]='TVRSU'; r[C['amount']]=primary['totalDist']
        out.append({'type':'TVRSU','txnId':rec['txnId'],'data':r})
        r=base(s,key,rec['empId']); r[C['rowId']]=line; line+=1
        r[C['earning']]='NVRSU'; r[C['amount']]=primary['totalWith']
        out.append({'type':'NVRSU','txnId':rec['txnId'],'data':r})
        for sr in rec['rows']:
            if (sr['regionTax'] in (0,None)) and (sr['incomeRegion'] in (0,None)): continue
            ln=line; line+=1
            r=base(s,key,rec['empId']); r[C['rowId']]=ln; r[C['deduction']]='W_SWW'
            r[C['amount']]=sr['regionTax']; r[C['stateAuthority']]=sr['stateCode']
            r[C['subRowId']]=1; r[C['relatedCalc']]='W_CWCGW'; r[C['value']]=sr['incomeRegion']
            out.append({'type':'W_SWW','txnId':rec['txnId'],'data':r})
            r=base(s,key,rec['empId']); r[C['rowId']]=ln; r[C['subRowId']]=2
            r[C['relatedCalc']]='W_TTWG'; r[C['value']]=sr['incomeRegion']
            out.append({'type':'sub','txnId':rec['txnId'],'data':r})
            r=base(s,key,rec['empId']); r[C['rowId']]=ln; r[C['subRowId']]=3
            r[C['relatedCalc']]='W_TXWG'; r[C['value']]=sr['incomeRegion']
            out.append({'type':'sub','txnId':rec['txnId'],'data':r})
    return out

def render_table(eib_rows):
    cls_map={'TVRSU':'tv','NVRSU':'nv','W_SWW':'ws','sub':'sb'}
    hdrs=''.join(f'<th>{l}</th>' for l,_ in DISPLAY)
    body=''; last=None
    for row in eib_rows:
        if row['txnId']!=last and last is not None:
            body+=f'<tr class="dv"><td colspan="{len(DISPLAY)}"></td></tr>'
        last=row['txnId']
        cls=cls_map.get(row['type'],'')
        cells=''.join(
            f'<td{"  class=\"num\"" if isinstance(row["data"][i],(int,float)) and not isinstance(row["data"][i],bool) else ""}>{fmt(row["data"][i])}</td>'
            for _,i in DISPLAY)
        body+=f'<tr class="{cls}">{cells}</tr>'
    return f'<div class="tbl-wrap"><table class="eib-tbl"><thead><tr>{hdrs}</tr></thead><tbody>{body}</tbody></table></div>'

def to_excel(eib_rows, s):
    buf=io.BytesIO(); wb=openpyxl.Workbook(); ws=wb.active
    ws.title='Payroll Off cycle Payment'
    hdr=empty()
    for i,h in enumerate(EIB_HDR): hdr[i]=h
    ws.append(hdr)
    for i,row in enumerate(eib_rows):
        d=list(row['data']); n=i+2
        d[C['paymentId']]=f'=CONCATENATE("Trailing2State",D{n},F{n})'
        ws.append(d)
    wb.save(buf); buf.seek(0); return buf.getvalue()

# ══════════════════════════════════════════════════════════
# UI
# ══════════════════════════════════════════════════════════

# ── Step 1: Upload ─────────────────────────────────────────
st.markdown('<div class="eib-card"><span class="badge">Step 1</span><h3>Load Output File</h3><p>Upload your output spreadsheet (Original or Deloitte sheet).</p></div>', unsafe_allow_html=True)
uploaded = st.file_uploader("Upload file", type=['xlsx','xls'], label_visibility='collapsed')

if uploaded is None:
    st.info("👆 Upload your output file above to get started.")

else:
    wb_src = openpyxl.load_workbook(uploaded, data_only=True)
    st.success(f"✅ **{uploaded.name}** loaded — sheets: {', '.join(wb_src.sheetnames)}")

    # ── Step 2: Select ─────────────────────────────────────
    st.markdown('<div class="eib-card"><span class="badge">Step 2</span><h3>Select Employee / Transaction</h3></div>', unsafe_allow_html=True)

    preferred = ['Original','Deloitte','QUESTIONS_']
    sheet_list = [s for s in preferred if s in wb_src.sheetnames] + \
                 [s for s in wb_src.sheetnames if s not in preferred]

    c1,c2,c3,c4 = st.columns([1,1,1,0.7])
    with c1: sheet_name = st.selectbox("Source Sheet", sheet_list, key='sheet')
    ws_src = wb_src[sheet_name]
    parsed = parse_sheet(ws_src)

    if not parsed:
        st.error("Could not find expected columns. Try a different sheet.")
    else:
        all_emps = sorted(set(r['empId'] for r in parsed))
        with c2: emp_id = st.selectbox("Employee ID", ['— choose —'] + all_emps, key='emp')
        if emp_id != '— choose —':
            all_txns = [r['txnId'] for r in parsed if r['empId'] == emp_id]
            with c3: txn_choice = st.selectbox("Transaction ID", ['All'] + all_txns, key='txn')
            with c4:
                st.write("")
                load = st.button("Load Data", type="primary", use_container_width=True)

            recs = [r for r in parsed if r['empId']==emp_id and
                    (txn_choice=='All' or r['txnId']==txn_choice)]

            # Source preview
            if load or st.session_state.get('loaded'):
                st.session_state['loaded'] = True
                preview=[]
                for rec in recs:
                    for sr in rec['rows']:
                        preview.append({'Emp ID':rec['empId'],'Txn ID':rec['txnId'],
                            'Name':rec['name'],'State':sr['stateCode'],
                            'Total Dist':sr['totalDist'],'EE Withholding':sr['totalWith'],
                            'Region Tax':sr['regionTax'],'Income for Region':sr['incomeRegion']})
                with st.expander("📋 Source data preview", expanded=False):
                    st.dataframe(pd.DataFrame(preview), use_container_width=True, hide_index=True)

                # ── Step 3: Settings ───────────────────────
                st.markdown('<div class="eib-card"><span class="badge">Step 3</span><h3>EIB Settings</h3><p>Stamped on every row to match your EIB template.</p></div>', unsafe_allow_html=True)
                s1,s2,s3,s4,s5,s6,s7 = st.columns(7)
                with s1: bid  = st.text_input("Batch ID",          "RSU May 26")
                with s2: pd_  = st.text_input("Payment Date",      "2026-05-01")
                with s3: prd  = st.text_input("Period Date",        "2026-05-01")
                with s4: pri  = st.text_input("Payment Priority",   "3")
                with s5: rcat = st.text_input("Run Category",       "RSU_USA")
                with s6: rtyp = st.text_input("Result Type",        "OnDemandPayment")
                with s7: rsn  = st.text_input("Reason",             "Stock")

                settings=dict(batchId=bid,paymentDate=pd_,periodDate=prd,
                              payPriority=pri,runCategory=rcat,resultType=rtyp,reason=rsn)

                # ── Step 4: Output ─────────────────────────
                eib_rows = gen(recs, settings)
                table = render_table(eib_rows)
                st.markdown(f'''
                <div class="eib-card">
                  <span class="badge">Step 4</span>
                  <h3>EIB Formatted Output <span class="row-badge">{len(eib_rows)} EIB rows</span></h3>
                  <div class="legend">
                    <div class="leg"><span class="dot dot-tv"></span> TVRSU – Total Distribution</div>
                    <div class="leg"><span class="dot dot-nv"></span> NVRSU – EE Withholding</div>
                    <div class="leg"><span class="dot dot-ws"></span> W_SWW – State Tax</div>
                    <div class="leg"><span class="dot dot-sb"></span> Sub-calc lines</div>
                  </div>
                  {table}
                </div>''', unsafe_allow_html=True)

                st.download_button("⬇️ Export to Excel (EIB format)",
                    data=to_excel(eib_rows, settings),
                    file_name="eib_output.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary")

                # ── Bulk ───────────────────────────────────
                st.markdown('<div class="eib-card"><span class="badge">Bulk</span><h3>Bulk Process All Employees</h3><p>Generate EIB rows for every employee in the selected sheet at once.</p></div>', unsafe_allow_html=True)
                if st.button("⚡ Generate All EIB Rows", type="secondary"):
                    all_eib = gen(parsed, settings)
                    all_tbl = render_table(all_eib)
                    st.success(f"✅ {len(all_eib)} EIB rows generated for {len(parsed)} transactions.")
                    st.markdown(f'<div class="eib-card">{all_tbl}</div>', unsafe_allow_html=True)
                    st.download_button("⬇️ Export All to Excel (EIB format)",
                        data=to_excel(all_eib, settings),
                        file_name="eib_bulk_output.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        type="primary", key="bulk_dl")
